#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""AutoDesign-LLM 전체 요약 파워포인트 생성 (python-pptx)"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # assets/ → 프로젝트 루트
A = os.path.join(ROOT, "assets")

# ---- 팔레트 (기술/자동차 느낌) ----
NAVY  = RGBColor(0x0B,0x25,0x45)
BLUE  = RGBColor(0x1A,0x73,0xE8)
TEAL  = RGBColor(0x00,0xA6,0xA6)
STEEL = RGBColor(0x1C,0x72,0x93)
LIGHT = RGBColor(0xF4,0xF6,0xF8)
WHITE = RGBColor(0xFF,0xFF,0xFF)
INK   = RGBColor(0x20,0x21,0x24)
MUTED = RGBColor(0x5F,0x63,0x68)
CARD  = RGBColor(0xEA,0xF1,0xFB)

KF = "Apple SD Gothic Neo"   # 한글 폰트

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s

def box(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in tf.margin_left, : pass
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    if isinstance(runs, str): runs = [[(runs, {})]]
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        if isinstance(para, tuple): para = [para]
        for t, st in para:
            r = p.add_run(); r.text = t
            f = r.font; f.name = KF
            f.size = Pt(st.get("sz", 14)); f.bold = st.get("b", False)
            f.italic = st.get("i", False)
            f.color.rgb = st.get("c", INK)
    return tb

def title_bar(s, num, t, dark=False):
    c = WHITE if dark else NAVY
    # 번호 원형 모티프
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(0.5), Inches(0.62), Inches(0.62))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE; circ.line.fill.background()
    circ.shadow.inherit = False
    ctf = circ.text_frame; ctf.word_wrap=False
    ctf.margin_left=0; ctf.margin_right=0; ctf.margin_top=0; ctf.margin_bottom=0
    cp = ctf.paragraphs[0]; cp.alignment=PP_ALIGN.CENTER
    cr = cp.add_run(); cr.text = num; cr.font.name=KF; cr.font.size=Pt(20); cr.font.bold=True; cr.font.color.rgb=WHITE
    text(s, 1.4, 0.5, 11.4, 0.85, [[(t, {"sz":30,"b":True,"c":c})]], anchor=MSO_ANCHOR.MIDDLE)

def img_fit(s, path, x, y, max_w, max_h, center_x=None):
    iw, ih = Image.open(path).size
    r = iw/ih
    w = max_w; h = w/r
    if h > max_h: h = max_h; w = h*r
    if center_x is not None: x = center_x - w/2
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h

def bullets(s, x, y, w, h, items, sz=15, gap=6, color=INK):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            txt, sub = it
            runs.append([("▸  ", {"sz":sz,"b":True,"c":BLUE}), (txt, {"sz":sz,"b":True,"c":color})])
            if sub: runs.append([("     "+sub, {"sz":sz-2,"c":MUTED})])
        else:
            runs.append([("▸  ", {"sz":sz,"b":True,"c":BLUE}), (it, {"sz":sz,"c":color})])
    text(s, x, y, w, h, runs, sp_after=gap)

def table(s, x, y, w, headers, rows, col_w, fs=11, hfs=12, rh=0.34):
    nr = len(rows)+1; nc = len(headers)
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rh*nr)).table
    for j, cw in enumerate(col_w): gt.columns[j].width = Inches(cw)
    gt.first_row = False; gt.horz_banding = False
    def cell(i,j,txt,b=False,fill=None,col=INK,size=fs,align=PP_ALIGN.LEFT):
        c = gt.cell(i,j)
        c.margin_left=Pt(5); c.margin_right=Pt(5); c.margin_top=Pt(2); c.margin_bottom=Pt(2)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fill is not None: c.fill.solid(); c.fill.fore_color.rgb=fill
        else: c.fill.solid(); c.fill.fore_color.rgb=WHITE
        tf=c.text_frame; tf.word_wrap=True
        p=tf.paragraphs[0]; p.alignment=align
        r=p.add_run(); r.text=txt; r.font.name=KF; r.font.size=Pt(size); r.font.bold=b
        r.font.color.rgb=col
    for j,hh in enumerate(headers):
        cell(0,j,hh,b=True,fill=NAVY,col=WHITE,size=hfs,align=PP_ALIGN.CENTER)
    for i,row in enumerate(rows):
        fill = LIGHT if i%2 else WHITE
        for j,val in enumerate(row):
            cell(i+1,j,val,fill=fill,align=PP_ALIGN.LEFT if j==len(row)-1 or j==1 else PP_ALIGN.CENTER if nc>3 else PP_ALIGN.LEFT)
    return gt

# ============================================================ SLIDES
# 1) 타이틀
s = slide(NAVY)
box(s, 0, 5.7, 13.333, 0.06, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 2.0, 11.5, 1.4, [[("AutoDesign-", {"sz":54,"b":True,"c":WHITE}),("LLM",{"sz":54,"b":True,"c":TEAL})]])
text(s, 0.95, 3.25, 11.5, 1.0, [[("자연어로 설계하고, 이론·표준으로 검증하는 자동차 부품 CAD 시스템", {"sz":22,"c":RGBColor(0xCA,0xDC,0xFC)})]])
text(s, 0.95, 6.0, 11.5, 0.6, [[("LLM × FreeCAD × FEM/CFD   ·   개발 기획 요약", {"sz":14,"c":TEAL})],
                               [("2026-06-09  ·  설계·기획 단계", {"sz":12,"c":RGBColor(0x9A,0xA0,0xA6)})]])

# 2) 개요 & 타깃
s = slide(WHITE); title_bar(s,"1","무엇을 만드는가")
text(s, 1.4, 1.5, 11.5, 0.9, [[("엔지니어가 ", {"sz":17}),("말로 요구사항을 설명하면", {"sz":17,"b":True,"c":BLUE}),
    (", 표준·이론에 근거해 사전 검토하고 물리적으로 검증된 자동차 부품 CAD를 만들어 주는 설계 가속 도구.", {"sz":17})]])
cards = [("설계 엔지니어","초기 설계·검증 반복을\n대폭 단축"),
         ("중소 부품사(Tier 2/3)","부족한 CAE 역량을\n저비용으로 보완"),
         ("컨설팅·스타트업","컨셉을 빠르게\n검토·제안")]
for i,(h,d) in enumerate(cards):
    x = 1.4 + i*3.85
    box(s, x, 2.7, 3.55, 2.0, fill=CARD, line=BLUE, lw=1.0)
    text(s, x+0.25, 2.95, 3.05, 0.6, [[(h,{"sz":16,"b":True,"c":NAVY})]])
    text(s, x+0.25, 3.6, 3.05, 1.0, [[(l,{"sz":13,"c":INK})] for l in d.split("\n")])
text(s, 1.4, 5.1, 11.5, 1.6, [
    [("핵심 가치 — ",{"sz":16,"b":True,"c":TEAL}),("단순 '말로 모델링'이 아니라 검증·사전검토 레이어가 차별점",{"sz":16,"b":True})],
    [("비목표 — 양산 최종 승인 도구가 아님(보조). 정식 CAE·실물시험을 대체하지 않음.",{"sz":13,"c":MUTED})],
], sp_after=8)

# 3) 핵심 차별점
s = slide(WHITE); title_bar(s,"2","핵심 차별점: 검증·사전검토 레이어")
text(s, 1.4, 1.5, 11.5, 0.6, [[("시중 '말로 3D 모델링' 도구는 형상 생성까지. 본 시스템은 자동차 도메인의 정량 검증을 자동화한다.",{"sz":15})]])
cols = [("기존 도구","· 자연어 → 형상 생성\n· 검증 없음\n· 도메인 무관","낮음", RGBColor(0xFC,0xE8,0xE6)),
        ("AutoDesign-LLM","· 형상 생성 +\n· 강도·모달·피로·DFM 검증\n· 표준·물성 근거 사전검토\n· 자기교정 루프","높음", RGBColor(0xE6,0xF4,0xEA))]
for i,(h,d,_,c) in enumerate(cols):
    x = 1.4 + i*5.9
    box(s, x, 2.3, 5.5, 3.5, fill=c, line=(BLUE if i else RGBColor(0xD9,0x30,0x25)), lw=1.2)
    text(s, x+0.35, 2.55, 4.9, 0.6, [[(h,{"sz":19,"b":True,"c":NAVY})]])
    text(s, x+0.35, 3.3, 4.9, 2.3, [[(l,{"sz":15})] for l in d.split("\n")], sp_after=8)

# 4) 전체 시스템 구성도
s = slide(WHITE); title_bar(s,"3","전체 시스템 구성도")
img_fit(s, os.path.join(A,"diagram_system_overview.png"), 0, 1.5, 11.0, 5.25, center_x=6.666)
text(s, 0, 6.95, 13.333, 0.4, [[("입력 → RAG 사전검토 → 형상 생성 → 다물리 검증 → 자기교정 루프 → 경량화 → 검증된 산출물",{"sz":12,"i":True,"c":MUTED})]], align=PP_ALIGN.CENTER)

# 5) 자동차 도메인 검증 레이어
s = slide(WHITE); title_bar(s,"4","자동차 도메인 검증 레이어 (핵심)")
table(s, 0.9, 1.55, 11.6,
      ["분류","검증 항목","방법","합격 기준(예)"],
      [["기하","솔리드 유효성","isValid()","닫힌 솔리드·자기교차 0"],
       ["구조","정적 강도","FEM(CalculiX)","안전계수 SF ≥ 목표"],
       ["진동","모달(NVH)","FEM 모달","1차 고유진동수 ≥ 임계"],
       ["내구","피로 수명","S-N + Miner","목표 사이클 초과"],
       ["제조","DFM","규칙+형상질의","구배·두께·R·언더컷 충족"],
       ["중량","경량화","위상최적화","SF 유지하며 질량 최소"]],
      [1.3,2.7,3.0,4.6], fs=13, hfs=13, rh=0.62)
text(s, 0.9, 6.35, 11.6, 0.8, [[("안전 원칙 — ",{"sz":13,"b":True,"c":RGBColor(0xD9,0x30,0x25)}),
    ("모든 물성·표준 수치는 검증된 출처로만. 최종 양산 판정은 정식 CAE·실물시험 필수.",{"sz":13,"c":INK})]])

# 6) 외형 보완 툴체인
s = slide(WHITE); title_bar(s,"5","외형(Exterior) 보완 — 오픈소스 툴체인")
text(s,1.4,1.5,11.5,0.5,[[("외형도 동일하게 '생성→검증→교정' 루프. 단, 검증 축이 FEM이 아니라 ",{"sz":14}),("공력(CFD)",{"sz":14,"b":True,"c":TEAL}),(".",{"sz":14})]])
steps = ["① AI 개념 생성\nTRELLIS / Hunyuan3D","② 리토폴로지\nBlender / Quadriflow",
         "③ 공력 검증 ★\nOpenFOAM (Cd·Cl)","④ 메시→NURBS\nFreeCAD RE / OpenNURBS"]
for i,st in enumerate(steps):
    x=1.4+i*2.95
    box(s,x,2.3,2.65,1.5,fill=CARD,line=BLUE,lw=1.0)
    text(s,x+0.15,2.45,2.35,1.2,[[(l,{"sz":13,"b": (j==0),"c":(NAVY if j==0 else INK)})] for j,l in enumerate(st.split("\n"))],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    if i<3:
        ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+2.68),Inches(2.85),Inches(0.24),Inches(0.4))
        ar.fill.solid();ar.fill.fore_color.rgb=BLUE;ar.line.fill.background();ar.shadow.inherit=False
bullets(s,1.4,4.3,11.4,2.6,[
    ("제조용 Class-A 서피스는 오픈소스로 아직 어렵다","→ Alias/ICEM 상용 영역. 목표는 '개념+공력+근사 서피스'"),
    ("AI 생성 메시는 '개념'이지 '설계'가 아니다","→ 리토폴로지·서피스화가 필수 다리"),
    ("외형 교정은 주관적 — 공력(Cd)은 정량, 미감은 사람 판단","→ 시스템은 공력·패키징·법규 정량 항목에 집중"),
], sz=14, gap=7)

# 7) 기획 로드맵
s = slide(WHITE); title_bar(s,"6","개발 로드맵 (기획자 관점)")
table(s, 0.7, 1.5, 12.0,
      ["Phase","테마","핵심 산출물 / Exit Gate"],
      [["P0","기반 구축","FreeCAD·CalculiX·LLM 추상화 / 골든 부품 생성 성공"],
       ["P1","형상 생성 MVP","자연어→브래킷 생성 + 기하 검증"],
       ["P2","구조 검증+교정","FEM 강도·SF 판정 / 미달→교정→통과 E2E"],
       ["P3","지식(RAG)+사전검토","RAG·물성DB / 리포트가 표준·물성 인용"],
       ["P4","다물리 검증+DFM","모달·피로·제조성 / 보고서(PDF) 산출"],
       ["P5","경량화","위상최적화 / 질량 ≥15%↓ & SF 유지"],
       ["P6","외형 툴체인","Blender/AI + OpenFOAM / 컨셉 외형+Cd"],
       ["P7","통합·제품화","UI·어셈블리 / 파일럿 PoC 완료"]],
      [1.0,2.6,8.4], fs=12.5, hfs=13, rh=0.6)

# 8) 마일스톤 & KPI
s = slide(WHITE); title_bar(s,"7","마일스톤 & 성공지표(KPI)")
box(s,0.9,1.55,5.7,5.1,fill=LIGHT,line=BLUE,lw=1.0)
text(s,1.2,1.75,5.1,0.5,[[("마일스톤",{"sz":18,"b":True,"c":NAVY})]])
bullets(s,1.2,2.4,5.2,4.0,[
    ("M1 (P2): '말로 검증된 브래킷 생성'","핵심 가설 증명"),
    ("M2 (P4): 풀 검증 + 검토보고서","베타"),
    ("M3 (P5): 경량화까지","차별화 완성"),
    ("M4 (P7): 파일럿 PoC","상용화 진입"),
], sz=14, gap=8)
box(s,6.9,1.55,5.6,5.1,fill=CARD,line=BLUE,lw=1.0)
text(s,7.2,1.75,5.0,0.5,[[("KPI",{"sz":18,"b":True,"c":NAVY})]])
bullets(s,7.2,2.4,5.0,4.0,[
    "검증 게이트 통과율 ≥ 70% (P4)",
    "골든 케이스 회귀 통과율 100% 유지",
    "요구사항→검증 CAD 시간 ≥ 50% 단축",
    "자기교정 루프 평균 ≤ 3회 수렴",
    "경량화: SF 유지 질량 ≥ 15%↓",
], sz=14, gap=10)

# 9) 개발 프로세스
s = slide(WHITE); title_bar(s,"8","개발 프로세스 (개발자 관점)")
items = [("인터페이스 우선 설계","LLMClient·Retriever·Validator 계약 뒤에 백엔드 은닉 → 세 버전이 코어 공유"),
         ("검증을 먼저, 생성을 나중에","검증 없으면 생성 품질 측정 불가 — 검증부터 구축"),
         ("골든 케이스 회귀","벤치마크 부품의 기대 결과 고정, LLM 비결정성은 '검증 통과·핵심 지표'로 판정"),
         ("샌드박스 실행 + 출처 추적","생성 Python 격리 실행, 물성·표준은 출처 메타 동반(임의생성 금지)"),
         ("CI/CD","lint→unit→integration(headless)→golden(허용오차)→배포, 무거운 시뮬은 nightly")]
y=1.6
for h,d in items:
    box(s,0.9,y,0.12,0.78,fill=BLUE,shape=MSO_SHAPE.RECTANGLE)
    text(s,1.2,y,11.4,0.8,[[(h,{"sz":15,"b":True,"c":NAVY})],[(d,{"sz":12.5,"c":INK})]],sp_after=2)
    y+=1.02

# 10) 세 버전 관계
s = slide(NAVY); title_bar(s,"9","세 가지 배포 버전",dark=True)
img_fit(s, os.path.join(A,"diagram_variants.png"), 0, 1.7, 12.4, 4.7, center_x=6.666)
text(s,0,6.6,13.333,0.5,[[("공통 코어 위에서 LLM·지식·연산 백엔드만 교체하는 하나의 진화 경로",{"sz":14,"i":True,"c":RGBColor(0xCA,0xDC,0xFC)})]],align=PP_ALIGN.CENTER)

# 11~13) 변형
def variant_slide(num,title,img,bg_note,pros,cons,note_color):
    s=slide(WHITE); title_bar(s,num,title)
    img_fit(s,os.path.join(A,img),6.7,1.55,6.3,4.4)
    bullets(s,0.9,1.7,5.6,2.6,[("강점",None)]+pros,sz=14,gap=6)
    text(s,0.9,1.7,5.6,0.4,[[("강점",{"sz":15,"b":True,"c":RGBColor(0x18,0x80,0x38)})]])
    # 재배치: 강점/약점 두 블록
    return s

# v1
s=slide(WHITE); title_bar(s,"10","v1 — 로컬 LLM (온프레미스)")
img_fit(s,os.path.join(A,"diagram_v1_local.png"),6.8,1.7,6.2,4.5)
text(s,0.9,1.6,5.7,0.4,[[("강점",{"sz":16,"b":True,"c":RGBColor(0x18,0x80,0x38)})]])
bullets(s,0.9,2.1,5.7,1.9,["IP 유출 0 (에어갭 가능)","API 비용 0 · 오프라인","데이터 주권"],sz=14,gap=6)
text(s,0.9,3.9,5.7,0.4,[[("약점 / 적합",{"sz":16,"b":True,"c":RGBColor(0xD9,0x30,0x25)})]])
bullets(s,0.9,4.4,5.7,2.3,["GPU 하드웨어 비용","도메인 지식 한계(RAG 전)","→ 방산·OEM 에어갭 환경"],sz=14,gap=6)

# v2
s=slide(WHITE); title_bar(s,"11","v2 — 로컬 + RAG 하이브리드 ★주력")
img_fit(s,os.path.join(A,"diagram_v2_hybrid.png"),6.8,1.7,6.2,4.5)
text(s,0.9,1.6,5.7,0.4,[[("강점",{"sz":16,"b":True,"c":RGBColor(0x18,0x80,0x38)})]])
bullets(s,0.9,2.1,5.7,1.9,["온프레미스 보안 유지","RAG로 도메인 정확도↑","사내 설계자산 활용"],sz=14,gap=6)
text(s,0.9,3.9,5.7,0.4,[[("약점 / 적합",{"sz":16,"b":True,"c":RGBColor(0xD9,0x30,0x25)})]])
bullets(s,0.9,4.4,5.7,2.3,["RAG 파이프라인·지식 큐레이션","→ 대다수 부품사·OEM 권장 기본값"],sz=14,gap=6)

# v3
s=slide(WHITE); title_bar(s,"12","v3 — 클라우드")
img_fit(s,os.path.join(A,"diagram_v3_cloud.png"),6.8,1.7,6.2,4.5)
text(s,0.9,1.6,5.7,0.4,[[("강점",{"sz":16,"b":True,"c":RGBColor(0x18,0x80,0x38)})]])
bullets(s,0.9,2.1,5.7,1.9,["최고 모델 품질(+비전)","무거운 시뮬 무한 확장","하드웨어 CAPEX 0 · 빠른 개발"],sz=14,gap=6)
text(s,0.9,3.9,5.7,0.4,[[("약점 / 적합",{"sz":16,"b":True,"c":RGBColor(0xD9,0x30,0x25)})]])
bullets(s,0.9,4.4,5.7,2.3,["데이터 외부 전송(IP 우려)","지속 OPEX","→ 비민감·대규모 시뮬·빠른 반복"],sz=14,gap=6)

# 14) 비교표
s=slide(WHITE); title_bar(s,"13","세 버전 비교")
table(s,0.7,1.55,12.0,
      ["기준","v1 로컬","v2 하이브리드 ★","v3 클라우드"],
      [["LLM 위치","로컬 GPU","로컬 GPU","클라우드 API"],
       ["지식(RAG)","없음/정적","로컬 RAG","매니지드 RAG"],
       ["도메인 정확도","중","높음","높음"],
       ["모델 품질","중","중","최고"],
       ["IP 보안","최고(에어갭)","높음","낮음"],
       ["무거운 시뮬 확장","제한","옵션(버스트)","최고"],
       ["비용 구조","CAPEX(GPU)","CAPEX+큐레이션","OPEX(토큰+연산)"],
       ["적합 고객","에어갭·방산","대다수 부품사·OEM","스타트업·대규모 시뮬"]],
      [2.6,2.7,3.4,3.3], fs=12, hfs=12.5, rh=0.58)

# 15) 리스크 & 안전
s=slide(WHITE); title_bar(s,"14","리스크 & 안전 원칙")
table(s,0.9,1.55,11.6,
      ["리스크","대응"],
      [["검증 신뢰성 부족 → 신뢰 상실","골든 케이스·엔지니어 검수, '보조 도구' 포지셔닝"],
       ["로컬 모델 코드생성 품질 부족","RAG 보강·코드특화 모델, 폴백으로 클라우드"],
       ["물성·표준 데이터 확보/라이선스","검증된 데이터 파트너십, 출처 관리"],
       ["무거운 시뮬 비용/시간","대리모델(surrogate)·클라우드 버스트"],
       ["부품 안전 책임(규제)","면책 고지·검수 게이트·감사 로그"]],
      [4.6,7.0], fs=13, hfs=13, rh=0.66)
text(s,0.9,6.4,11.6,0.7,[[("안전 원칙 — 시스템은 엔지니어를 ",{"sz":14,"b":True}),("보조",{"sz":14,"b":True,"c":BLUE}),
    ("한다. 최종 양산 판정은 정식 CAE·실물시험을 거쳐야 한다.",{"sz":14,"b":True})]])

# 16) 마무리
s=slide(NAVY)
box(s,0,5.6,13.333,0.06,fill=TEAL,shape=MSO_SHAPE.RECTANGLE)
text(s,0.9,2.1,11.5,1.0,[[("다음 단계",{"sz":40,"b":True,"c":WHITE})]])
bullets(s,1.0,3.3,11.3,2.2,[
    "물성·표준 데이터 소스 확보·라이선스 조사 (리드타임 최장)",
    "Phase 0 코어 스캐폴딩 — LLM 추상화 + 엔진 브래킷 골든 케이스",
    "주력 v2(하이브리드) 기준으로 RAG 지식베이스 설계",
], sz=18, gap=12, color=RGBColor(0xCA,0xDC,0xFC))
text(s,0.95,6.0,11.5,0.6,[[("AutoDesign-LLM  ·  2026-06-09  ·  설계·기획 단계",{"sz":13,"c":TEAL})]])

out = os.path.join(ROOT, "AutoDesign-LLM_요약.pptx")
prs.save(out)
print("SAVED", out, "slides=", len(prs.slides._sldIdLst))
