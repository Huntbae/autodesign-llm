#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""마크다운 문서 → Word(.docx) 변환 (다이어그램 이미지 삽입 포함)"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # assets/ → 프로젝트 루트
ASSETS = os.path.join(ROOT, "assets")
OUT = os.path.join(ROOT, "word")
os.makedirs(OUT, exist_ok=True)

# 파일별 (제목 포함 텍스트, 이미지파일, 캡션) — 해당 헤딩 직후 이미지 삽입
IMG_MAP = {
    "DESIGN.md": [("시스템 아키텍처", "diagram_system_overview.png", "그림. 전체 시스템 구성도")],
    "00-INDEX.md": [("세 버전의 관계", "diagram_variants.png", "그림. 세 버전의 진화 관계")],
    "01-ROADMAP.md": [("배포 버전 출시 전략", "diagram_variants.png", "그림. 세 버전의 진화 관계")],
    "03-VARIANT-LOCAL.md": [("아키텍처", "diagram_v1_local.png", "그림. v1 로컬 LLM 아키텍처")],
    "04-VARIANT-HYBRID.md": [("아키텍처", "diagram_v2_hybrid.png", "그림. v2 하이브리드 아키텍처")],
    "05-VARIANT-CLOUD.md": [("아키텍처", "diagram_v3_cloud.png", "그림. v3 클라우드 아키텍처")],
}

BOX = set("┌┐└┘│─┄►◄▼▲├┤┬┴┼╔╗╚╝║═")

def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def add_inline(par, text):
    # **bold**, `code`, [t](url) 처리
    token = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        t = m.group(0)
        if t.startswith("**"):
            r = par.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`"):
            r = par.add_run(t[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:  # link
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", t)
            r = par.add_run(lm.group(1)); r.font.color.rgb = RGBColor(0x1A,0x73,0xE8); r.underline = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])

def add_code(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.1); pf.space_before = Pt(4); pf.space_after = Pt(4)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"; run.font.size = Pt(7.5)
    # 연한 배경
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),"F5F5F5")
    pPr.append(shd)

def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if not re.match(r"^[\s:|-]+$", rows[i].strip().strip("|"))]
    if not cells: return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(cells):
        tr = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cell = tr[ci]; cell.text = ""
            par = cell.paragraphs[0]
            add_inline(par, txt)
            for run in par.runs:
                run.font.size = Pt(9)
                if ri == 0: run.bold = True
            if ri == 0:
                set_cell_shading(cell, "1A73E8")
                for run in par.runs: run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

def insert_image(doc, img, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(ASSETS, img), width=Inches(6.3))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption); r.italic = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5F,0x63,0x68)

def convert(md_path):
    name = os.path.basename(md_path)
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    doc = Document()
    doc.styles["Normal"].font.name = "Malgun Gothic"
    doc.styles["Normal"].font.size = Pt(10.5)
    imgs = IMG_MAP.get(name, [])
    i = 0
    while i < len(lines):
        line = lines[i]
        # 코드펜스
        if line.strip().startswith("```"):
            block = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i]); i += 1
            i += 1
            add_code(doc, block)
            continue
        # 테이블
        if line.strip().startswith("|") and i+1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i+1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(doc, rows)
            continue
        # 헤딩
        h = re.match(r"^(#{1,4})\s+(.*)$", line)
        if h:
            level = len(h.group(1)); txt = h.group(2).strip()
            if level == 1:
                p = doc.add_heading(txt, level=0)
            else:
                p = doc.add_heading("", level=min(level-1, 4))
                add_inline(p, txt)
            # 이미지 삽입 트리거
            for trig, img, cap in imgs:
                if trig in txt:
                    insert_image(doc, img, cap)
            i += 1
            continue
        # 수평선
        if re.match(r"^---+$", line.strip()):
            i += 1; continue
        # 블록쿼트
        if line.strip().startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
            r_txt = line.strip().lstrip(">").strip()
            add_inline(p, r_txt)
            for r in p.runs: r.italic = True; r.font.color.rgb = RGBColor(0x44,0x44,0x44)
            i += 1; continue
        # 리스트
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if lm:
            indent = len(lm.group(1))
            style = "List Bullet" if lm.group(2) in ("-","*") else "List Number"
            p = doc.add_paragraph(style=style)
            if indent >= 2: p.paragraph_format.left_indent = Inches(0.5)
            add_inline(p, lm.group(3))
            i += 1; continue
        # 빈 줄
        if not line.strip():
            i += 1; continue
        # 일반 문단
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1
    out_path = os.path.join(OUT, name.replace(".md", ".docx"))
    doc.save(out_path)
    print("wrote", out_path)

if __name__ == "__main__":
    targets = ["DESIGN.md", "PROGRESS.md",
               "docs/00-INDEX.md", "docs/01-ROADMAP.md", "docs/02-DEV-PROCESS.md",
               "docs/03-VARIANT-LOCAL.md", "docs/04-VARIANT-HYBRID.md", "docs/05-VARIANT-CLOUD.md",
               "docs/06-REAL-LLM-BUILD.md", "docs/07-HERMES-LOCAL.md", "docs/08-DEV-SETUP.md"]
    for t in targets:
        convert(os.path.join(ROOT, t))
    print("ALL DOCX DONE")
